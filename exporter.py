"""Plain-text to PDF / Word, and web page text extraction."""
from __future__ import annotations

import io
import re
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

NAVY = HexColor("#0F2C4C")
INK = HexColor("#1C2433")
MUTED = HexColor("#5C6B7A")

HEADING_NAMES = {
    "executive summary",
    "core competencies & technical skillset",
    "professional experience",
    "education & certifications",
}

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
)


def fetch_web_plain_text(url: str, timeout: int = 15, max_bytes: int = 2_000_000) -> str:
    parsed = urlparse(url.strip())
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        raise ValueError("Only http and https URLs are allowed.")
    response = requests.get(
        url.strip(),
        timeout=timeout,
        headers={"User-Agent": USER_AGENT, "Accept": "text/html,application/xhtml+xml"},
        stream=True,
    )
    response.raise_for_status()
    content_type = (response.headers.get("Content-Type") or "").lower()
    if "html" not in content_type and "text/plain" not in content_type and "xml" not in content_type:
        raise ValueError("That URL is not an HTML or text page.")
    chunks = []
    total = 0
    for chunk in response.iter_content(chunk_size=16384):
        total += len(chunk)
        if total > max_bytes:
            raise ValueError("Page is larger than 2 MB.")
        chunks.append(chunk)
    raw = b"".join(chunks)
    if "text/plain" in content_type:
        return raw.decode(response.encoding or "utf-8", errors="replace").strip()
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "iframe", "nav", "footer", "header"]):
        tag.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    body = soup.get_text("\n", strip=True)
    body = re.sub(r"\n{3,}", "\n\n", body).strip()
    if title:
        return f"{title}\n\n{body}"
    return body


def _is_heading(line: str) -> bool:
    return line.strip().lower() in HEADING_NAMES


def _is_name_line(line: str, index: int) -> bool:
    stripped = line.strip()
    return index == 0 or (
        stripped.upper() == stripped
        and 2 <= len(stripped.split()) <= 4
        and stripped.replace(" ", "").isalpha()
    )


def text_to_pdf_bytes(text: str, title: str = "Document") -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
        title=title,
        author="Resume Studio",
    )
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "Name", parent=styles["Normal"], fontName="Times-Bold", fontSize=16,
        textColor=NAVY, leading=20, spaceAfter=2, alignment=1,
    )
    role_style = ParagraphStyle(
        "Role", parent=styles["Normal"], fontName="Times-Italic", fontSize=11,
        textColor=MUTED, leading=14, spaceAfter=6, alignment=1,
    )
    contact_style = ParagraphStyle(
        "Contact", parent=styles["Normal"], fontName="Times-Roman", fontSize=9,
        textColor=INK, leading=12, spaceAfter=10, alignment=1,
    )
    heading_style = ParagraphStyle(
        "Head", parent=styles["Normal"], fontName="Times-Bold", fontSize=12,
        textColor=NAVY, leading=16, spaceBefore=12, spaceAfter=6,
        borderPadding=2,
    )
    body_style = ParagraphStyle(
        "Body", parent=styles["Normal"], fontName="Times-Roman", fontSize=10.5,
        textColor=INK, leading=14, spaceAfter=7, alignment=4,
    )
    label_style = ParagraphStyle(
        "Label", parent=styles["Normal"], fontName="Times-Bold", fontSize=10.5,
        textColor=INK, leading=14, spaceAfter=2,
    )

    story = []
    lines = [ln.rstrip() for ln in text.replace("\r\n", "\n").split("\n")]
    i = 0
    seen_name = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        escaped = (
            line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        )
        if seen_name < 2 and line.upper() == line and 2 <= len(line.split()) <= 4 and line.replace(" ", "").isalpha():
            if seen_name == 1:
                story.append(Spacer(1, 16))
            story.append(Paragraph(escaped, name_style))
            seen_name += 1
            i += 1
            continue
        if line.lower().startswith("senior it program manager"):
            story.append(Paragraph(escaped, role_style))
            i += 1
            continue
        if line.startswith("Location:"):
            story.append(Paragraph(escaped.replace("|", " | "), contact_style))
            i += 1
            continue
        if _is_heading(line):
            story.append(Paragraph(escaped.upper(), heading_style))
            i += 1
            continue
        if line.endswith(":") and len(line) < 80 and not line.startswith("Dear"):
            story.append(Paragraph(f"<b>{escaped}</b>", label_style))
            i += 1
            continue
        story.append(Paragraph(escaped, body_style))
        i += 1

    doc.build(story)
    return buffer.getvalue()


def text_to_docx_bytes(text: str, title: str = "Document") -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles["Normal"]
    styles.font.name = "Times New Roman"
    styles.font.size = Pt(11)
    styles.font.color.rgb = RGBColor(0x1C, 0x24, 0x33)

    seen_name = 0
    for index, raw in enumerate(text.replace("\r\n", "\n").split("\n")):
        line = raw.strip()
        if not line:
            continue
        if seen_name < 2 and line.strip().isupper() and 2 <= len(line.strip().split()) <= 4:
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            run.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x0F, 0x2C, 0x4C)
            run.font.name = "Times New Roman"
            seen_name += 1
            continue
        if line.lower().startswith("senior it program manager"):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            run.italic = True
            run.font.size = Pt(11)
            run.font.name = "Times New Roman"
            continue
        if line.startswith("Location:"):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line)
            run.font.size = Pt(9)
            run.font.name = "Times New Roman"
            continue
        if _is_heading(line):
            p = document.add_paragraph()
            run = p.add_run(line.upper())
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x0F, 0x2C, 0x4C)
            run.font.name = "Times New Roman"
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(line)
        if line.endswith(":") and len(line) < 80 and not line.startswith("Dear"):
            run.bold = True
        run.font.size = Pt(11)
        run.font.name = "Times New Roman"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
